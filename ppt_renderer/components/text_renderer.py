"""
ppt_renderer/components/text_renderer.py
========================================
Text component renderer.

Handles plain text, list items, labels, and generic content boxes.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN

from ppt_renderer.components.coordinates import convert_bounds
from ppt_renderer.components.placeholder_resolver import resolve_placeholder
from backend.design_system import design_language
from backend.design_system.theme_loader import get_current_theme


def render(
    component_specification,
    presentation,
    slide,
    content: dict[str, Any],
    *,
    layout_context: dict[str, Any] | None = None,
) -> None:
    """Render a text-based component."""
    theme = get_current_theme()
    left, top, width, height = convert_bounds(component_specification, presentation)
    pattern_id = (layout_context or {}).get("pattern_id")
    value = resolve_placeholder(component_specification.placeholder, content)

    if isinstance(value, dict):
        text = value.get("text", value.get("title", ""))
    elif isinstance(value, list):
        text = "\n".join(str(item) for item in value)
    else:
        text = str(value)

    context_alignment = design_language.get_alignment(
        component_specification.type, pattern_id=pattern_id
    )
    alignment = _to_pp_align(context_alignment)
    emphasis = design_language.get_emphasis(
        "primary" if component_specification.type == "label" else "secondary",
        pattern_id=pattern_id,
    )

    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text or ""
    p.alignment = alignment
    p.font.name = theme.font("body")
    p.font.size = theme.size("body")
    p.font.bold = emphasis.get("bold", component_specification.type == "label")
    p.font.color.rgb = theme.color(emphasis.get("color_role", "ink"))


def _to_pp_align(value: str):
    """Map a design-language alignment name to a python-pptx constant."""
    mapping = {
        "center": PP_ALIGN.CENTER,
        "left": PP_ALIGN.LEFT,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }
    return mapping.get(str(value).lower(), PP_ALIGN.LEFT)
